import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips

PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_ROOT / "docs" / "TaskFlow_API_Project_Guide_ZH_v4.docx"

BODY_FONT = "Microsoft YaHei"
CODE_FONT = "Consolas"
TEXT_COLOR = "000000"
MUTED_COLOR = "595959"
BORDER_COLOR = "D9D9D9"
HEADER_FILL = "1F4E79"
ALT_ROW_FILL = "F2F6FA"
CODE_FILL = "F4F4F4"


def set_run_font(run, name: str) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def set_style_font(style, name: str) -> None:
    style.font.name = name
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def remove_paragraph_borders(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def remove_style_borders(style) -> None:
    p_pr = style._element.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), BORDER_COLOR)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    table_header = OxmlElement("w:tblHeader")
    table_header.set(qn("w:val"), "true")
    tr_pr.append(table_header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def ensure_table_child(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def set_table_width(parent, tag: str, width_dxa: int) -> None:
    width = ensure_table_child(parent, tag)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(int(width_dxa)))


def column_widths_from_weights(weights, total_width_dxa: int) -> list[int]:
    total_weight = float(sum(weights))
    widths = [
        int(round(total_width_dxa * (weight / total_weight))) for weight in weights
    ]
    widths[-1] += total_width_dxa - sum(widths)
    return widths


def apply_table_geometry(table, column_widths_dxa) -> None:
    widths = [int(width) for width in column_widths_dxa]
    table_width_dxa = sum(widths)
    cell_margin_dxa = 120

    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    table_properties = table._tbl.tblPr
    set_table_width(table_properties, "w:tblW", table_width_dxa)
    set_table_width(table_properties, "w:tblInd", cell_margin_dxa)

    layout = ensure_table_child(table_properties, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_column = OxmlElement("w:gridCol")
        grid_column.set(qn("w:w"), str(width))
        grid.append(grid_column)

    for index, width in enumerate(widths):
        table.columns[index].width = Twips(width)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[index]
            cell.width = Twips(width)
            set_table_width(cell._tc.get_or_add_tcPr(), "w:tcW", width)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("TaskFlow API 项目教程    ")
    set_run_font(run, BODY_FONT)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED_COLOR)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    add_page_number(section.footer.paragraphs[0])

    normal = document.styles["Normal"]
    set_style_font(normal, BODY_FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT_COLOR)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    title = document.styles["Title"]
    set_style_font(title, BODY_FONT)
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(TEXT_COLOR)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(10)
    remove_style_borders(title)

    heading_config = {
        "Heading 1": (18, 14, 8),
        "Heading 2": (14, 12, 6),
        "Heading 3": (12, 9, 5),
    }
    for style_name, (size, before, after) in heading_config.items():
        style = document.styles[style_name]
        set_style_font(style, BODY_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(TEXT_COLOR)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        remove_style_borders(style)

    document.core_properties.title = "TaskFlow API Python 后端项目教程"
    document.core_properties.subject = "项目说明、七天操作教程、代码走读、简历与面试指南"
    document.core_properties.author = "TaskFlow API 项目文档"
    document.core_properties.comments = "Generated from the local TaskFlow API project."


def add_inline_text(paragraph, text: str) -> None:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    token_pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    for token in token_pattern.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, CODE_FONT)
            run.font.size = Pt(9)
        else:
            run = paragraph.add_run(token)
            set_run_font(run, BODY_FONT)


def add_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    add_inline_text(paragraph, text)
    paragraph.paragraph_format.space_after = Pt(6)


def add_list_item(document: Document, text: str, number: str | None = None) -> None:
    paragraph = document.add_paragraph()
    prefix = f"{number} " if number else "• "
    paragraph.add_run(prefix)
    add_inline_text(paragraph, text)
    paragraph.paragraph_format.left_indent = Inches(0.28)
    paragraph.paragraph_format.first_line_indent = Inches(-0.22)
    paragraph.paragraph_format.space_after = Pt(3)


def add_code_block(document: Document, code: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.05)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.0
    shade_paragraph(paragraph, CODE_FILL)
    lines = code.rstrip().splitlines()
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        set_run_font(run, CODE_FONT)
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(TEXT_COLOR)


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return

    column_count = len(rows[0])
    table = document.add_table(rows=1, cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False

    available_width_dxa = 9720
    weights = []
    for column_index in range(column_count):
        samples = [len(row[column_index]) for row in rows[: min(len(rows), 8)]]
        weights.append(max(4, sum(samples) / max(1, len(samples))))
    widths = column_widths_from_weights(weights, available_width_dxa)

    header_cells = table.rows[0].cells
    for index, value in enumerate(rows[0]):
        cell = header_cells[index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(value)
        set_run_font(run, BODY_FONT)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)
        shade_cell(cell, HEADER_FILL)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_repeat_table_header(table.rows[0])

    for row_index, row_data in enumerate(rows[1:], start=1):
        cells = table.add_row().cells
        fill = ALT_ROW_FILL if row_index % 2 == 0 else "FFFFFF"
        for column_index, value in enumerate(row_data):
            cell = cells[column_index]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if column_count <= 4 and len(value) <= 14
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            add_inline_text(paragraph, value)
            for run in paragraph.runs:
                if run.font.size is None:
                    run.font.size = Pt(9)
            shade_cell(cell, fill)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for row in table.rows:
        set_row_cant_split(row)
    apply_table_geometry(table, widths)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def parse_markdown(markdown_text: str) -> list[tuple]:
    lines = markdown_text.splitlines()
    blocks: list[tuple] = []
    index = 0

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        if stripped.startswith("```"):
            code_lines = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            blocks.append(("code", "\n".join(code_lines)))
            index += 1
            continue

        heading_match = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading_match:
            blocks.append(("heading", len(heading_match.group(1)), heading_match.group(2)))
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines):
            separator = lines[index + 1].strip()
            if re.match(r"^\|[\s:-]+\|", separator):
                table_rows = []
                header = [cell.strip() for cell in stripped.strip("|").split("|")]
                table_rows.append(header)
                index += 2
                while index < len(lines) and lines[index].strip().startswith("|"):
                    row = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                    table_rows.append(row)
                    index += 1
                blocks.append(("table", table_rows))
                continue

        bullet_match = re.match(r"^[-*]\s+(.+)$", stripped)
        if bullet_match:
            blocks.append(("bullet", bullet_match.group(1)))
            index += 1
            continue

        number_match = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if number_match:
            blocks.append(("number", number_match.group(1) + ".", number_match.group(2)))
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate:
                break
            if candidate.startswith(("#", "```", "|", "- ", "* ")):
                break
            if re.match(r"^\d+\.\s+", candidate):
                break
            paragraph_lines.append(candidate)
            index += 1
        blocks.append(("paragraph", " ".join(paragraph_lines)))

    return blocks


def render_blocks(document: Document, blocks: list[tuple]) -> None:
    for block in blocks:
        block_type = block[0]
        if block_type == "heading":
            _, level, text = block
            if level == 1:
                level = 2
            document.add_heading(text, level=level)
        elif block_type == "paragraph":
            add_paragraph(document, block[1])
        elif block_type == "bullet":
            add_list_item(document, block[1])
        elif block_type == "number":
            add_list_item(document, block[2], block[1])
        elif block_type == "code":
            add_code_block(document, block[1])
        elif block_type == "table":
            add_table(document, block[1])


def add_cover(document: Document) -> None:
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(72)
    run = title.add_run("TaskFlow API Python 后端项目教程")
    set_run_font(run, BODY_FONT)
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor.from_string(TEXT_COLOR)
    remove_paragraph_borders(title)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(34)
    run = subtitle.add_run("项目说明、七天操作教程、代码走读、简历与面试指南")
    set_run_font(run, BODY_FONT)
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string(MUTED_COLOR)

    cover_rows = [
        ["项目名称", "TaskFlow API"],
        ["项目类型", "FastAPI 多用户任务管理后端"],
        ["适用读者", "准备 Python 后端简历项目的学习者"],
        ["本地环境", "Windows、Python 3.14.6"],
        ["验证状态", "12 项测试通过，测试覆盖率 94%"],
        ["项目位置", r"<你的项目目录>\taskflow-api"],
    ]
    add_table(document, cover_rows)

    generated = document.add_paragraph()
    generated.alignment = WD_ALIGN_PARAGRAPH.CENTER
    generated.paragraph_format.space_before = Pt(24)
    run = generated.add_run(f"生成日期：{date.today().isoformat()}")
    set_run_font(run, BODY_FONT)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MUTED_COLOR)
    document.add_page_break()


def read_after_title(path: Path) -> str:
    text = path.read_text(encoding="utf-8").strip()
    lines = text.splitlines()
    if lines and lines[0].startswith("# "):
        lines = lines[1:]
    return "\n".join(lines).strip()


def build_document() -> Path:
    document = Document()
    configure_document(document)
    add_cover(document)

    front_matter = """
## 文档说明

这份文档把 TaskFlow API 的项目说明、操作教程、接口速查、代码走读和简历准备内容整理为一本连续的学习手册。你可以先按“七天学习路线”建立整体认识，再从第一天开始逐项执行命令。每个阶段都有明确检查点，不需要一次看懂所有源代码。

这份文档面向首次完成完整 Python 后端项目的学习者。最终目标不是把代码复制到电脑里，而是能够独立启动、验证、解释并继续扩展项目。完成文档中的全部步骤后，你可以把项目写进简历，并在面试中回答鉴权、权限隔离、数据库设计、自动化测试和部署相关问题。

## 项目成果与验证结果

项目已在当前电脑完成真实运行验证。以下结果来自最终代码，不包含尚未执行的 Docker 构建。

| 验证项目 | 结果 | 说明 |
| --- | --- | --- |
| 自动化测试 | 12 项通过 | 覆盖注册、登录、越权、筛选、统计和级联删除 |
| 代码质量 | Ruff 通过 | 无导入、未使用变量和格式问题 |
| 覆盖率 | 94% | 基于 pytest-cov 统计 |
| HTTP 冒烟测试 | 通过 | 注册、登录、项目、任务、统计和删除流程完整返回正常 |
| 浏览器操作台 | 通过 | 已手工完成登录、创建项目、创建任务、标记完成和查看统计 |
| Docker | 未执行 | 当前电脑未安装 Docker，配置已经准备完成 |

## 项目概览

TaskFlow API 是一个多用户项目与任务管理后端。使用者可以注册账号、登录、创建项目、在项目中维护任务，并通过状态、优先级和关键词筛选任务。统计接口返回任务总数、各状态数量、逾期数量和完成率。

项目重点不在简单的增删改查，而在完整的后端工程链路。身份验证使用 OAuth2 Password Flow 和 JWT，密码通过 scrypt 加盐哈希保存。每个项目和任务都执行资源所有权检查，用户不能读取或修改其他用户的数据。项目还提供 `/demo` 浏览器操作台，便于不熟悉 Swagger 的使用者直接验收登录、项目、任务和统计流程。开发环境默认使用 SQLite，Docker 环境可以切换 PostgreSQL。

## 技术栈

| 层级 | 技术 | 作用 |
| --- | --- | --- |
| Web 框架 | FastAPI | 路由、依赖注入、参数校验和 OpenAPI 文档 |
| ORM | SQLAlchemy 2.0 | 模型定义、查询和事务管理 |
| 数据校验 | Pydantic 2 | 请求体和响应体校验 |
| 数据库 | SQLite、PostgreSQL | 本地零配置开发和容器化演示 |
| 鉴权 | OAuth2、JWT | 登录和接口身份验证 |
| 密码存储 | hashlib.scrypt | 加盐密码哈希 |
| 测试 | pytest、TestClient | 接口级自动化测试 |
| 质量与部署 | Ruff、GitHub Actions、Docker | 代码检查、持续集成和容器化 |

## 架构说明

请求先进入 FastAPI 路由。公共依赖负责解析 JWT、查询当前用户，并确认目标项目属于当前用户。路由随后调用 SQLAlchemy 会话完成数据库操作。数据库层通过环境变量选择 SQLite 或 PostgreSQL。

```text
HTTP Request
    |
FastAPI Router
    |
CurrentUser / OwnedProject dependency
    |
SQLAlchemy Session
    |
SQLite or PostgreSQL
```

这种设计把身份和权限判断集中在依赖层。路由代码不需要重复编写“这个资源是不是当前用户的”逻辑，减少遗漏权限检查的风险。

## 核心接口

| 方法 | 路径 | 作用 |
| --- | --- | --- |
| POST | `/api/v1/auth/register` | 注册用户 |
| POST | `/api/v1/auth/login` | 登录并获取 JWT |
| GET | `/api/v1/auth/me` | 读取当前用户 |
| GET、POST | `/api/v1/projects` | 项目列表和创建项目 |
| GET、PATCH、DELETE | `/api/v1/projects/{id}` | 项目详情、修改和删除 |
| GET | `/api/v1/projects/{id}/stats` | 项目任务统计 |
| GET、POST | `/api/v1/projects/{id}/tasks` | 任务列表和创建任务 |
| GET、PATCH、DELETE | `/api/v1/tasks/{id}` | 任务详情、修改和删除 |

## 七天学习路线

| 天数 | 主要任务 | 完成标志 |
| --- | --- | --- |
| 第 1 天 | 创建环境并启动服务 | 浏览器能打开 `/demo` 操作台 |
| 第 2 天 | 在 `/demo` 操作台中走通完整业务 | 能登录、创建项目、创建任务、标记完成并查看统计 |
| 第 3 天 | 理解目录和请求流程 | 能说明一次创建任务请求经过哪些层 |
| 第 4 天 | 理解鉴权和用户隔离 | 能解释 JWT 和 scoped resource 检查 |
| 第 5 天 | 运行测试并制造一次失败 | 能定位断言失败并恢复测试 |
| 第 6 天 | 使用 Docker 和 PostgreSQL | 容器启动后数据能够持久化 |
| 第 7 天 | 发布到 GitHub | CI 通过且没有提交敏感文件 |
"""
    render_blocks(document, parse_markdown(front_matter))

    document.add_page_break()
    document.add_heading("完整操作教程", level=1)
    tutorial = read_after_title(PROJECT_ROOT / "TUTORIAL.md")
    render_blocks(document, parse_markdown(tutorial))

    document.add_page_break()
    document.add_heading("接口速查", level=1)
    api_reference = """
## 通用请求要求

除注册、登录和健康检查以外，所有接口都需要在请求头中携带：

```text
Authorization: Bearer <access_token>
```

列表接口使用 `page` 和 `page_size` 分页。`page` 从 1 开始，`page_size` 范围为 1 到 100。项目和任务搜索使用 `search` 参数。

## 项目接口示例

创建项目：

```json
{
  "name": "TaskFlow Launch",
  "description": "完成第一版后端并准备部署"
}
```

任务列表筛选：

```text
GET /api/v1/projects/1/tasks?status=in_progress&priority=high&page=1&page_size=20
```

任务状态值：

| 值 | 含义 |
| --- | --- |
| `todo` | 待处理 |
| `in_progress` | 进行中 |
| `done` | 已完成 |
| `cancelled` | 已取消 |

任务优先级值：

| 值 | 含义 |
| --- | --- |
| `low` | 低 |
| `medium` | 中 |
| `high` | 高 |
| `urgent` | 紧急 |

## 统计响应示例

```json
{
  "total": 3,
  "todo": 0,
  "in_progress": 1,
  "done": 1,
  "cancelled": 1,
  "overdue": 0,
  "completion_rate": 33.33
}
```
"""
    render_blocks(document, parse_markdown(api_reference))

    document.add_page_break()
    document.add_heading("代码走读与面试讲解", level=1)
    code_walkthrough = read_after_title(PROJECT_ROOT / "docs" / "CODE_WALKTHROUGH.md")
    render_blocks(document, parse_markdown(code_walkthrough))

    document.add_page_break()
    document.add_heading("简历写法与面试准备", level=1)
    resume_guide = read_after_title(PROJECT_ROOT / "docs" / "RESUME_GUIDE.md")
    render_blocks(document, parse_markdown(resume_guide))

    document.add_page_break()
    document.add_heading("完成标准", level=1)
    completion = r"""
你只有在能够独立完成以下任务时，才应该把项目写入简历并声称自己掌握相关技术：

- 从空终端启动服务并打开 `/demo` 操作台。
- 创建两个用户并证明他们无法访问对方的资源。
- 解释密码为什么不保存明文。
- 解释 JWT 中 `sub` 和 `exp` 的含义。
- 运行 Ruff 和 pytest，并处理一次测试失败。
- 说明 SQLite 与 PostgreSQL 的用途差异。
- 把项目推送到 GitHub，并确认 CI 通过。
- 用 60 秒说明项目背景、方案、难点、验证结果和不足。

## 后续扩展顺序

| 顺序 | 扩展内容 | 主要收益 |
| --- | --- | --- |
| 1 | 任务标签与多标签筛选 | 练习多对多关系和复杂查询 |
| 2 | 项目成员和角色权限 | 练习 RBAC 与权限边界 |
| 3 | Refresh token 和登出 | 练习令牌生命周期管理 |
| 4 | Alembic 数据库迁移 | 补足生产数据库演进能力 |
| 5 | Redis 统计缓存 | 练习缓存设计和一致性 |
| 6 | 云部署、HTTPS 和监控 | 形成可公开演示的完整项目 |

## 最终检查

在发送简历前，重新执行：

```powershell
cd "C:\\path\\to\\taskflow-api"
.\.venv\Scripts\python.exe -m ruff check .
.\.venv\Scripts\python.exe -m pytest --cov=taskflow --cov-report=term-missing
.\.venv\Scripts\python.exe -m uvicorn taskflow.main:app --reload
```

服务启动后，再执行一次真实 HTTP 冒烟测试：

```powershell
.\.venv\Scripts\python.exe scripts\smoke_test.py
```

所有检查通过后，再更新 GitHub 仓库和简历描述。
"""
    render_blocks(document, parse_markdown(completion))

    document.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    output = build_document()
    print(output)
