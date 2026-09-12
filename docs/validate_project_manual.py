from collections import Counter
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = PROJECT_ROOT / "docs" / "TaskFlow_API_Project_Guide_ZH_v4.docx"


def validate() -> None:
    assert DOCX_PATH.exists(), f"Missing DOCX: {DOCX_PATH}"

    with ZipFile(DOCX_PATH) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")

    assert "<w:cantSplit" in document_xml
    assert "<w:tblHeader" in document_xml
    assert "<w:pBdr" not in document_xml

    document = Document(DOCX_PATH)
    style_counts = Counter(paragraph.style.name for paragraph in document.paragraphs)
    assert style_counts["Title"] == 1
    assert style_counts["Heading 1"] == 5
    assert style_counts["Heading 2"] >= 25
    assert style_counts["Heading 3"] >= 25
    assert len(document.tables) >= 8

    empty_cells = 0
    oversized_cells = 0
    for table in document.tables:
        for row in table.rows:
            assert row._tr.find(qn("w:trPr")).find(qn("w:cantSplit")) is not None
            for cell in row.cells:
                value = cell.text.strip()
                if not value:
                    empty_cells += 1
                if len(value) > 500:
                    oversized_cells += 1

    assert empty_cells == 0
    assert oversized_cells == 0

    paragraph_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    document_text = f"{paragraph_text}\n{table_text}"
    assert "TODO" not in paragraph_text
    assert "REPLACE_ME" not in paragraph_text
    assert "TaskFlow API Python 后端项目教程" in paragraph_text
    assert "12 项测试通过" in document_text
    assert "http://127.0.0.1:8000/demo" in document_text
    assert "USER ISOLATION PASSED" in document_text

    print(f"DOCX: {DOCX_PATH}")
    print(f"Paragraphs: {len(document.paragraphs)}")
    print(f"Tables: {len(document.tables)}")
    print(f"Heading 1: {style_counts['Heading 1']}")
    print(f"Heading 2: {style_counts['Heading 2']}")
    print(f"Heading 3: {style_counts['Heading 3']}")
    print("Structural validation passed")


if __name__ == "__main__":
    validate()
